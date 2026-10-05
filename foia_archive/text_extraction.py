"""Text extraction and bounded OCR for archived FOIA documents."""
from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Optional

import pypdfium2 as pdfium
import pytesseract
from docx import Document
from PIL import Image
from pypdf import PdfReader
from pytesseract.pytesseract import TesseractNotFoundError


TEXT_LIKE_TYPES = {
    "txt",
    "csv",
    "json",
    "xml",
    "rtf",
    "eml",
}
OCR_IMAGE_TYPES = {
    "png",
    "jpg",
    "jpeg",
    "tif",
    "tiff",
}


@dataclass(frozen=True)
class OCRSettings:
    enabled: bool = False
    languages: str = "eng"
    dpi: int = 200
    max_pages_per_document: int = 50
    page_timeout_seconds: float = 45.0
    max_image_megapixels: float = 20.0

    @classmethod
    def from_mapping(cls, value: Optional[Mapping[str, object]]) -> "OCRSettings":
        data = value or {}

        enabled_raw = data.get("enabled", False)
        if isinstance(enabled_raw, str):
            enabled = enabled_raw.strip().lower() in {"1", "true", "yes", "on"}
        else:
            enabled = bool(enabled_raw)

        def _int(name: str, default: int, minimum: int, maximum: int) -> int:
            try:
                parsed = int(data.get(name, default))
            except (TypeError, ValueError):
                parsed = default
            return max(minimum, min(maximum, parsed))

        def _float(
            name: str,
            default: float,
            minimum: float,
            maximum: float,
        ) -> float:
            try:
                parsed = float(data.get(name, default))
            except (TypeError, ValueError):
                parsed = default
            return max(minimum, min(maximum, parsed))

        languages = str(data.get("languages") or "eng").strip() or "eng"

        return cls(
            enabled=enabled,
            languages=languages,
            dpi=_int("dpi", 200, 72, 400),
            max_pages_per_document=_int(
                "max_pages_per_document",
                50,
                1,
                500,
            ),
            page_timeout_seconds=_float(
                "page_timeout_seconds",
                45.0,
                1.0,
                300.0,
            ),
            max_image_megapixels=_float(
                "max_image_megapixels",
                20.0,
                1.0,
                100.0,
            ),
        )


@dataclass(frozen=True)
class TextExtractionResult:
    status: str
    text: str = ""
    error: Optional[str] = None
    method: Optional[str] = None
    character_count: int = 0
    truncated: bool = False


class OCRUnavailable(RuntimeError):
    pass


class OCRPageFailure(RuntimeError):
    pass


def _cap_text(
    text: str,
    max_chars: int,
    *,
    method: Optional[str],
    error: Optional[str] = None,
    force_truncated: bool = False,
) -> TextExtractionResult:
    normalized = text.replace("\x00", " ").strip()
    if not normalized:
        return TextExtractionResult(
            status="empty",
            error=error or "No extractable text was found.",
            method=method,
        )

    truncated = force_truncated or len(normalized) > max_chars
    indexed = normalized[:max_chars]
    return TextExtractionResult(
        status="indexed_truncated" if truncated else "indexed",
        text=indexed,
        error=error,
        method=method,
        character_count=len(indexed),
        truncated=truncated,
    )


def _bounded_image(image: Image.Image, max_megapixels: float) -> Image.Image:
    max_pixels = max(1, int(max_megapixels * 1_000_000))
    current_pixels = max(1, image.width * image.height)
    if current_pixels <= max_pixels:
        return image

    scale = math.sqrt(max_pixels / current_pixels)
    width = max(1, int(image.width * scale))
    height = max(1, int(image.height * scale))
    resampling = getattr(Image, "Resampling", Image).LANCZOS
    return image.resize((width, height), resampling)


def _ocr_image(image: Image.Image, settings: OCRSettings) -> str:
    prepared = _bounded_image(
        image.convert("L"),
        settings.max_image_megapixels,
    )
    try:
        return pytesseract.image_to_string(
            prepared,
            lang=settings.languages,
            timeout=settings.page_timeout_seconds,
        )
    except TesseractNotFoundError as exc:
        raise OCRUnavailable(
            "Tesseract OCR executable is not installed or not on PATH."
        ) from exc
    except RuntimeError as exc:
        raise OCRPageFailure(
            f"Tesseract OCR timed out or failed: {exc}"
        ) from exc
    except pytesseract.TesseractError as exc:
        raise OCRPageFailure(
            f"Tesseract OCR failed: {exc}"
        ) from exc


def _render_pdf_page(
    pdf: pdfium.PdfDocument,
    page_index: int,
    settings: OCRSettings,
) -> Image.Image:
    page = pdf[page_index]
    bitmap = None
    try:
        bitmap = page.render(scale=settings.dpi / 72.0)
        return bitmap.to_pil().convert("RGB").copy()
    finally:
        if bitmap is not None:
            bitmap.close()
        page.close()


def _extract_pdf(
    path: Path,
    max_chars: int,
    ocr: OCRSettings,
) -> TextExtractionResult:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return TextExtractionResult(
                status="extraction_failed",
                error="Encrypted PDF could not be opened without a password.",
            )

    pdfium_doc: Optional[pdfium.PdfDocument] = None
    parts: list[str] = []
    native_used = False
    ocr_used = False
    blank_pages = 0
    ocr_attempted = 0
    ocr_failures: list[str] = []
    ocr_unavailable = False
    incomplete = False
    current_chars = 0

    try:
        for page_index, page in enumerate(reader.pages):
            native_text = (page.extract_text() or "").strip()
            if native_text:
                native_used = True
                parts.append(native_text)
                current_chars += len(native_text) + 1
            else:
                blank_pages += 1
                if not ocr.enabled:
                    continue
                if ocr_unavailable:
                    incomplete = True
                    continue
                if ocr_attempted >= ocr.max_pages_per_document:
                    incomplete = True
                    continue

                ocr_attempted += 1
                try:
                    if pdfium_doc is None:
                        pdfium_doc = pdfium.PdfDocument(str(path))
                    image = _render_pdf_page(
                        pdfium_doc,
                        page_index,
                        ocr,
                    )
                    try:
                        ocr_text = _ocr_image(image, ocr).strip()
                    finally:
                        image.close()
                    if ocr_text:
                        ocr_used = True
                        parts.append(ocr_text)
                        current_chars += len(ocr_text) + 1
                except OCRUnavailable as exc:
                    ocr_unavailable = True
                    incomplete = True
                    ocr_failures.append(str(exc))
                except OCRPageFailure as exc:
                    incomplete = True
                    ocr_failures.append(
                        f"page {page_index + 1}: {exc}"
                    )

            if current_chars > max_chars:
                incomplete = True
                break
    finally:
        if pdfium_doc is not None:
            pdfium_doc.close()

    if native_used and ocr_used:
        method = "mixed"
    elif ocr_used:
        method = "ocr"
    elif native_used:
        method = "native_text"
    else:
        method = None

    if parts:
        errors: list[str] = []
        if ocr_unavailable and blank_pages:
            errors.append(
                "Some image-only pages were not OCRed because Tesseract is unavailable."
            )
        elif ocr_failures:
            errors.append(
                f"{len(ocr_failures)} OCR page(s) failed; first error: "
                f"{ocr_failures[0]}"
            )
        if (
            ocr.enabled
            and blank_pages > ocr.max_pages_per_document
        ):
            errors.append(
                "OCR page safety limit reached; remaining image-only pages were skipped."
            )
        return _cap_text(
            "\n".join(parts),
            max_chars,
            method=method,
            error=" ".join(errors) or None,
            force_truncated=incomplete,
        )

    if blank_pages and ocr.enabled:
        if ocr_unavailable:
            return TextExtractionResult(
                status="ocr_unavailable",
                error=ocr_failures[0] if ocr_failures else (
                    "Tesseract OCR is unavailable."
                ),
                method="ocr",
            )
        if ocr_failures:
            return TextExtractionResult(
                status="ocr_failed",
                error=ocr_failures[0],
                method="ocr",
            )
        return TextExtractionResult(
            status="empty",
            error="OCR completed but no searchable text was recognized.",
            method="ocr",
        )

    return TextExtractionResult(
        status="empty",
        error=(
            "No extractable PDF text was found; OCR is disabled."
            if blank_pages
            else "No extractable PDF text was found."
        ),
    )


def _extract_docx(path: Path, max_chars: int) -> TextExtractionResult:
    document = Document(str(path))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))

    return _cap_text(
        "\n".join(parts),
        max_chars,
        method="native_text",
    )


def _extract_text_like(path: Path, max_chars: int) -> TextExtractionResult:
    # Read a bounded byte window so indexing cannot create an unbounded memory
    # spike even when the archive download ceiling is much larger.
    byte_limit = max(4096, max_chars * 4)
    with path.open("rb") as source:
        raw = source.read(byte_limit + 1)

    bytes_truncated = len(raw) > byte_limit
    raw = raw[:byte_limit]
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")

    return _cap_text(
        text,
        max_chars,
        method="native_text",
        force_truncated=bytes_truncated,
    )


def _extract_image(
    path: Path,
    max_chars: int,
    ocr: OCRSettings,
) -> TextExtractionResult:
    if not ocr.enabled:
        return TextExtractionResult(
            status="unsupported",
            error="Image OCR is disabled.",
        )

    parts: list[str] = []
    failures: list[str] = []
    incomplete = False

    try:
        with Image.open(path) as image:
            frame_count = int(getattr(image, "n_frames", 1) or 1)
            pages_to_process = min(
                frame_count,
                ocr.max_pages_per_document,
            )
            if frame_count > pages_to_process:
                incomplete = True

            for frame_index in range(pages_to_process):
                try:
                    image.seek(frame_index)
                    frame = image.copy()
                    try:
                        text = _ocr_image(frame, ocr).strip()
                    finally:
                        frame.close()
                    if text:
                        parts.append(text)
                except OCRUnavailable as exc:
                    return TextExtractionResult(
                        status="ocr_unavailable",
                        error=str(exc),
                        method="ocr",
                    )
                except OCRPageFailure as exc:
                    incomplete = True
                    failures.append(
                        f"page {frame_index + 1}: {exc}"
                    )
    except OCRUnavailable as exc:
        return TextExtractionResult(
            status="ocr_unavailable",
            error=str(exc),
            method="ocr",
        )

    if parts:
        error = None
        if failures:
            error = (
                f"{len(failures)} OCR page(s) failed; first error: "
                f"{failures[0]}"
            )
        elif incomplete:
            error = (
                "OCR page safety limit reached; remaining image pages were skipped."
            )
        return _cap_text(
            "\n".join(parts),
            max_chars,
            method="ocr",
            error=error,
            force_truncated=incomplete,
        )

    if failures:
        return TextExtractionResult(
            status="ocr_failed",
            error=failures[0],
            method="ocr",
        )

    return TextExtractionResult(
        status="empty",
        error="OCR completed but no searchable text was recognized.",
        method="ocr",
    )


def extract_document_text(
    path: Path | str,
    file_type: Optional[str],
    *,
    max_chars: int = 5_000_000,
    ocr: Optional[OCRSettings] = None,
) -> TextExtractionResult:
    """Extract searchable text without affecting archive success."""
    source = Path(path)
    kind = (file_type or source.suffix.lstrip(".")).strip().lower()
    max_chars = max(1_000, int(max_chars))
    ocr_settings = ocr or OCRSettings()

    try:
        if kind == "pdf":
            return _extract_pdf(source, max_chars, ocr_settings)
        if kind == "docx":
            return _extract_docx(source, max_chars)
        if kind in TEXT_LIKE_TYPES:
            return _extract_text_like(source, max_chars)
        if kind in OCR_IMAGE_TYPES:
            return _extract_image(source, max_chars, ocr_settings)
        return TextExtractionResult(
            status="unsupported",
            error=f"No text extractor is configured for file type: {kind or 'unknown'}",
        )
    except TesseractNotFoundError as exc:
        return TextExtractionResult(
            status="ocr_unavailable",
            error=f"TesseractNotFoundError: {exc}"[:2000],
            method="ocr",
        )
    except Exception as exc:
        return TextExtractionResult(
            status="extraction_failed",
            error=f"{type(exc).__name__}: {exc}"[:2000],
        )
