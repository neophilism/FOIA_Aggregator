import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from foia_archive.storage import (
    get_connection,
    get_document_text,
    get_schema_version,
    init_db,
    insert_document,
    query_documents,
    update_download_failure,
    upsert_document_text,
)
from foia_archive.text_extraction import extract_document_text
from foia_archive.text_index import reindex_downloaded_documents
from foia_archive.utils import Config


class FullTextSearchTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        self.db_path = self.root / "archive.db"
        self.files_dir = self.root / "files"
        init_db(self.db_path, self.files_dir)
        self.conn = get_connection(self.db_path)

    def tearDown(self):
        self.conn.close()
        self.tempdir.cleanup()

    def _insert(self, title="Unrelated title", filename="record.pdf"):
        return insert_document(
            self.conn,
            url=f"https://example.gov/{filename}",
            title=title,
            file_type=Path(filename).suffix.lstrip(".") or "pdf",
            filename=filename,
            agency_id=None,
            office_id=None,
            reading_room_id=None,
            discovered_at="2026-10-05T00:00:00",
        )

    def test_schema_migration_creates_fts5_tables(self):
        self.assertEqual(get_schema_version(self.conn), 6)
        names = {
            row["name"]
            for row in self.conn.execute(
                "SELECT name FROM sqlite_master WHERE name IN ('document_text', 'document_fts')"
            ).fetchall()
        }
        self.assertEqual(names, {"document_text", "document_fts"})

    def test_body_only_search_match_is_returned(self):
        doc_id = self._insert()
        upsert_document_text(
            self.conn,
            doc_id,
            body="The hidden phrase is constitutional restoration protocol.",
            extraction_status="indexed",
            extracted_at=datetime.now(timezone.utc).isoformat(),
        )

        rows = query_documents(
            self.conn,
            title_query="constitutional restoration",
        )

        self.assertEqual([row["id"] for row in rows], [doc_id])

    def test_plain_search_punctuation_cannot_break_fts_query(self):
        doc_id = self._insert()
        upsert_document_text(
            self.conn,
            doc_id,
            body="alpha beta gamma",
            extraction_status="indexed",
            extracted_at=datetime.now(timezone.utc).isoformat(),
        )

        rows = query_documents(
            self.conn,
            title_query='"alpha" + beta:(gamma)',
        )

        self.assertEqual([row["id"] for row in rows], [doc_id])

    def test_reindex_replaces_stale_fts_body(self):
        doc_id = self._insert()
        now = datetime.now(timezone.utc).isoformat()
        upsert_document_text(
            self.conn,
            doc_id,
            body="old searchable phrase",
            extraction_status="indexed",
            extracted_at=now,
        )
        upsert_document_text(
            self.conn,
            doc_id,
            body="new searchable phrase",
            extraction_status="indexed",
            extracted_at=now,
        )

        old_rows = query_documents(self.conn, title_query="old searchable")
        new_rows = query_documents(self.conn, title_query="new searchable")

        self.assertEqual(old_rows, [])
        self.assertEqual([row["id"] for row in new_rows], [doc_id])

    def test_download_failure_clears_stale_full_text(self):
        doc_id = self._insert()
        now = datetime.now(timezone.utc).isoformat()
        upsert_document_text(
            self.conn,
            doc_id,
            body="stale body should disappear",
            extraction_status="indexed",
            extracted_at=now,
        )

        update_download_failure(
            self.conn,
            doc_id,
            "http_error",
            "redownload failed",
            now,
        )

        self.assertEqual(
            query_documents(self.conn, title_query="stale body"),
            [],
        )
        self.assertIsNone(get_document_text(self.conn, doc_id))

    def test_nonindexed_status_removes_previous_fts_body(self):
        doc_id = self._insert()
        now = datetime.now(timezone.utc).isoformat()
        upsert_document_text(
            self.conn,
            doc_id,
            body="formerly searchable",
            extraction_status="indexed",
            extracted_at=now,
        )
        upsert_document_text(
            self.conn,
            doc_id,
            body="",
            extraction_status="extraction_failed",
            extraction_error="broken",
            extracted_at=now,
        )

        self.assertEqual(
            query_documents(self.conn, title_query="formerly searchable"),
            [],
        )
        row = get_document_text(self.conn, doc_id)
        self.assertEqual(row["extraction_status"], "extraction_failed")
        self.assertEqual(row["extraction_error"], "broken")


class TextExtractionTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)

    def tearDown(self):
        self.tempdir.cleanup()

    def test_extracts_plain_text_and_truncates_at_limit(self):
        path = self.root / "record.txt"
        path.write_text("alpha " * 1000, encoding="utf-8")

        result = extract_document_text(
            path,
            "txt",
            max_chars=1000,
        )

        self.assertEqual(result.status, "indexed_truncated")
        self.assertTrue(result.truncated)
        self.assertEqual(len(result.text), 1000)
        self.assertIn("alpha", result.text)

    def test_extracts_docx_paragraphs_and_tables(self):
        path = self.root / "record.docx"
        document = Document()
        document.add_paragraph("Paragraph searchable text")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Table Alpha"
        table.cell(0, 1).text = "Table Beta"
        document.save(path)

        result = extract_document_text(path, "docx")

        self.assertEqual(result.status, "indexed")
        self.assertIn("Paragraph searchable text", result.text)
        self.assertIn("Table Alpha", result.text)
        self.assertIn("Table Beta", result.text)

    def test_blank_pdf_is_marked_empty_for_future_ocr(self):
        path = self.root / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with path.open("wb") as target:
            writer.write(target)

        result = extract_document_text(path, "pdf")

        self.assertEqual(result.status, "empty")
        self.assertIn("OCR", result.error)

    def test_unsupported_binary_type_is_explicit(self):
        path = self.root / "record.zip"
        path.write_bytes(b"not really a zip")

        result = extract_document_text(path, "zip")

        self.assertEqual(result.status, "unsupported")
        self.assertIn("zip", result.error)


class ReindexExistingLocalArchiveTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        self.db_path = self.root / "archive.db"
        self.files_dir = self.root / "files"
        init_db(self.db_path, self.files_dir)

        self.config = Config(
            {
                "storage": {
                    "backend": "local",
                    "db_path": str(self.db_path),
                    "files_dir": str(self.files_dir),
                },
                "search": {
                    "max_indexed_chars_per_document": 100000,
                },
            }
        )

    def tearDown(self):
        self.tempdir.cleanup()

    def test_reindexes_existing_downloaded_local_file_without_recrawl(self):
        archive_path = self.files_dir / "existing" / "memo.txt"
        archive_path.parent.mkdir(parents=True)
        archive_path.write_text(
            "Unique body phrase from an already archived release.",
            encoding="utf-8",
        )

        conn = get_connection(self.db_path)
        try:
            doc_id = insert_document(
                conn,
                url="https://example.gov/memo.txt",
                title="Generic Memo",
                file_type="txt",
                filename="memo.txt",
                agency_id=None,
                office_id=None,
                reading_room_id=None,
                discovered_at="2026-10-05T00:00:00",
            )
            conn.execute(
                """
                UPDATE documents
                SET download_status = 'downloaded',
                    local_path = ?,
                    storage_backend = 'local',
                    storage_key = ?
                WHERE id = ?
                """,
                ("existing/memo.txt", "existing/memo.txt", doc_id),
            )
            conn.commit()
        finally:
            conn.close()

        summary = reindex_downloaded_documents(self.config)

        self.assertEqual(summary.attempted, 1)
        self.assertEqual(summary.indexed, 1)
        self.assertEqual(summary.failed, 0)

        conn = get_connection(self.db_path)
        try:
            rows = query_documents(
                conn,
                title_query="already archived release",
            )
            text_row = get_document_text(conn, doc_id)
        finally:
            conn.close()

        self.assertEqual([row["id"] for row in rows], [doc_id])
        self.assertEqual(text_row["extraction_status"], "indexed")


if __name__ == "__main__":
    unittest.main()
